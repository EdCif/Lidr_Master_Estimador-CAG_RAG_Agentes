"""Convierte ``results.csv`` en un Excel que se pueda leer (``results.xlsx``).

Para qué sirve
--------------
El CSV del stress test es para programas: 42 columnas con nombres técnicos,
comas como separador (Excel en español espera punto y coma) y 1/0 en vez de
sí/no. Este módulo hace una versión para personas, con tres hojas:

* **Turnos**: una fila por turno con las columnas importantes, en castellano,
  ✅/❌, segundos en vez de milisegundos, filtros y las filas con error en
  rojo. Dos columnas cuentan QUÉ se ha olvidado (de la conversación y del
  PDF) con palabras, no con códigos.
* **Leyenda**: qué significa cada columna y de dónde sale.
* **Hoja de corrección**: los tres guiones, con el dato que planta cada turno
  y hasta cuándo se comprueba. Sirve para cruzar "qué se ha olvidado" con
  "qué se dijo".

El CSV sigue siendo el entregable que pide el enunciado; esto es una ayuda
para revisarlo. El runner (``run.py``) lo genera solo al terminar, al lado
del CSV. También se puede lanzar a mano, desde la carpeta estimador-cag:

    uv run python -m evals.stress.to_excel --csv evals/stress/results.csv
"""

from __future__ import annotations

import argparse
import csv
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from evals.stress.fixtures.build_pdfs import PLANTED_FACTS
from evals.stress.scenarios import SCENARIOS, TURNS_PER_SCENARIO, Fact

SCENARIO_NAMES = {"growing": "Crece", "pivot": "Cambio de rumbo", "contradiction": "Contradicción"}
PLACE_NAMES = {"metadata": "memoria", "history": "historial", "reply": "respuesta"}
_PLACE_WITH_ARTICLE = {"metadata": "la memoria", "history": "el historial", "reply": "la respuesta"}

# Colores: los del panel de Streamlit (dorado y crema) y un rojo suave.
_HEADER_FILL = PatternFill("solid", fgColor="8C6D46")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ERROR_FILL = PatternFill("solid", fgColor="F8D7DA")
_TITLE_FONT = Font(bold=True, size=13)
_WRAP = Alignment(wrap_text=True, vertical="top")

# Etiqueta del hecho -> el hecho, por escenario (para decir qué se olvidó).
_FACTS: dict[str, dict[str, Fact]] = {
    name: {fact.label: fact for turn in scenario.turns for fact in turn.facts_to_remember}
    for name, scenario in SCENARIOS.items()
}
_PDF_FACTS = {fact.label: fact for fact in PLANTED_FACTS}


# ---------------------------------------------------------------------------
# 1. Cómo se escribe cada valor
# ---------------------------------------------------------------------------

def _number(value: str) -> float | None:
    return float(value) if value not in ("", None) else None


def _seconds(value: str) -> float | None:
    number = _number(value)
    return None if number is None else round(number / 1000, 1)


def _yes_no(value: str) -> str:
    return {"1": "✅", "0": "❌"}.get(value, "")


def _pairs(details: str) -> list[tuple[str, str]]:
    """"budget_80k=1;team_5=0" -> [("budget_80k", "1"), ("team_5", "0")]."""
    return [tuple(item.split("=", 1)) for item in (details or "").split(";") if "=" in item]


def _forgotten_in_conversation(row: dict[str, str]) -> str:
    """Los hechos de la hoja de corrección que han fallado, con palabras."""
    if not row.get("memory_drift_checked") or row["memory_drift_checked"] == "0":
        return ""
    facts = _FACTS.get(row["scenario"], {})
    failed = []
    for label, passed in _pairs(row.get("memory_drift_details", "")):
        if passed == "1":
            continue
        fact = facts.get(label)
        if fact is None:
            failed.append(label)
        elif fact.expected:
            failed.append(f"«{fact.text}»")
        else:
            places = " y ".join(_PLACE_WITH_ARTICLE[place] for place in fact.where)
            failed.append(f"«{fact.text}» sigue en {places} (debía desaparecer)")
    return "; ".join(failed) or "—"


def _pdf_in_reply(row: dict[str, str]) -> str:
    """Qué datos del PDF aparecen en la respuesta de este turno."""
    if row.get("attachment_kb") in ("", "0"):
        return "sin PDF"
    pairs = _pairs(row.get("attachment_recall_details", ""))
    if not pairs:
        return ""
    missing = [f"«{_PDF_FACTS[label].fact}»" for label, passed in pairs if passed != "1" and label in _PDF_FACTS]
    if not missing:
        return f"✅ los {len(pairs)}"
    if len(missing) == len(pairs):
        return "❌ ninguno"
    return "falta " + ", ".join(missing)


def _outcome(row: dict[str, str]) -> str:
    if row.get("outcome") == "done":
        return "✅"
    return f"❌ {(row.get('error_detail') or row.get('outcome') or '')[:150]}"


@dataclass(frozen=True)
class Column:
    """Una columna de la hoja Turnos: título, ancho, cómo se calcula y qué significa."""

    title: str
    width: int
    value: Callable[[dict[str, str]], Any]
    meaning: str
    source: str
    number_format: str | None = None


COLUMNS = [
    Column("Escenario", 15, lambda r: SCENARIO_NAMES.get(r["scenario"], r["scenario"]),
           "Qué conversación: el proyecto que crece, el que cambia de rumbo o el que se contradice.",
           "scenario"),
    Column("PDF (KB)", 9, lambda r: int(r["attachment_kb"]),
           "Tamaño del PDF adjuntado en el turno 1, en miles de caracteres de texto. 0 = sin PDF.",
           "attachment_kb"),
    Column("Repetición", 10, lambda r: int(r["repeat"]),
           "Número de repetición: la misma conversación se lanza varias veces para ver cuánto varía.",
           "repeat"),
    Column("Turno", 7, lambda r: int(r["turn_index"]),
           "Número de mensaje dentro de la conversación (1 a 20).", "turn_index"),
    Column("Resultado", 30, _outcome,
           "✅ si el turno terminó bien; ❌ y el motivo si falló (el turno no se guardó).",
           "outcome, error_detail"),
    Column("Tokens entrada", 12, lambda r: _number(r["tokens_in"]),
           "Lo que leyó el modelo: system prompt (con la memoria), historial de la ventana, "
           "mensaje y PDF. Un token son unos 4 caracteres.", "tokens_in", "#,##0"),
    Column("Tokens salida", 12, lambda r: _number(r["tokens_out"]),
           "Lo que escribió el modelo, incluido su razonamiento interno (que no se ve pero se paga).",
           "tokens_out", "#,##0"),
    Column("Coste turno ($)", 13, lambda r: _number(r["cost_usd"]),
           "Lo que costó este turno, con los precios de app/services/model_costs.py.",
           "cost_usd", "0.0000"),
    Column("Coste acumulado ($)", 14, lambda r: _number(r["cumulative_cost_usd"]),
           "Lo que lleva gastado la conversación hasta este turno, incluido.",
           "cumulative_cost_usd", "0.0000"),
    Column("Latencia (s)", 11, lambda r: _seconds(r["latency_ms"]),
           "Lo que tardó la API en completar el turno (extraer el PDF, llamar al modelo y "
           "actualizar la memoria), sin contar la red.", "latency_ms", "0.0"),
    Column("Primer texto (s)", 12, lambda r: _seconds(r["ttft_ms"]),
           "Desde que se llama al modelo hasta que llega el primer trozo de texto: lo que tarda en "
           "«leer» lo que le mandamos.", "ttft_ms", "0.0"),
    Column("Mensajes en ventana", 12, lambda r: _number(r["messages_in_window"]),
           "Mensajes del historial tras el turno (máximo 12: 6 turnos de pregunta y respuesta).",
           "messages_in_window", "0"),
    Column("Hechos en memoria", 11, lambda r: _number(r["anchors_count"]),
           "Datos guardados en la ficha del proyecto: nombre, equipo, cada tecnología y alcance.",
           "anchors_count", "0"),
    Column("¿Latencia OK?", 10, lambda r: _yes_no(r["latency_budget_passed"]),
           "¿El turno tardó menos que el presupuesto de latencia (20 s por defecto)?",
           "latency_budget_passed"),
    Column("¿Coste OK?", 9, lambda r: _yes_no(r["cost_budget_passed"]),
           "¿El turno costó menos que el presupuesto de coste (0,03 $ por defecto)?",
           "cost_budget_passed"),
    Column("Recuerda (%)", 10, lambda r: _number(r["memory_drift_score"]),
           "De los datos de la hoja de corrección que tocaba comprobar en este turno, qué parte "
           "sigue viva (en la memoria, el historial o la respuesta, según el dato).",
           "memory_drift_score", "0%"),
    Column("Olvidado o mal", 40, _forgotten_in_conversation,
           "Qué datos de la conversación han fallado en este turno: los que se han perdido, o los "
           "que siguen cuando debían haber desaparecido. «—» = ninguno.",
           "memory_drift_details + evals/stress/scenarios.py"),
    Column("Datos del PDF en la respuesta", 26, _pdf_in_reply,
           "Cuáles de los tres datos escondidos en el PDF (SAP, 12.000 reservas, ISO 27001) usa "
           "el modelo en su respuesta de este turno.",
           "attachment_recall_details + evals/stress/fixtures/build_pdfs.py"),
]


# ---------------------------------------------------------------------------
# 2. Las tres hojas
# ---------------------------------------------------------------------------

def _header(sheet, titles: list[str], widths: list[int]) -> None:
    sheet.append(titles)
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _turns_sheet(workbook: Workbook, rows: list[dict[str, str]]) -> None:
    sheet = workbook.active
    sheet.title = "Turnos"
    _header(sheet, [column.title for column in COLUMNS], [column.width for column in COLUMNS])
    for row in rows:
        sheet.append([column.value(row) for column in COLUMNS])
        excel_row = sheet.max_row
        for index, column in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=excel_row, column=index)
            if column.number_format:
                cell.number_format = column.number_format
            if row.get("outcome") != "done":
                cell.fill = _ERROR_FILL
    sheet.auto_filter.ref = sheet.dimensions


def _legend_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Leyenda")
    _header(sheet, ["Columna", "Qué significa", "Columna(s) del CSV"], [26, 90, 40])
    for column in COLUMNS:
        sheet.append([column.title, column.meaning, column.source])
        for cell in sheet[sheet.max_row]:
            cell.alignment = _WRAP
    sheet.append([])
    sheet.append(["Escenarios"])
    sheet.cell(row=sheet.max_row, column=1).font = _TITLE_FONT
    for name, scenario in SCENARIOS.items():
        sheet.append([SCENARIO_NAMES[name], scenario.description, name])
        for cell in sheet[sheet.max_row]:
            cell.alignment = _WRAP


def _answer_key_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Hoja de corrección")
    _header(
        sheet,
        ["Escenario", "Turno", "Lo que dice el cliente", "Dato a comprobar", "Dónde se busca",
         "Debe…", "Se comprueba tras los turnos"],
        [15, 7, 70, 24, 22, 18, 16],
    )
    for name, scenario in SCENARIOS.items():
        for turn in scenario.turns:
            facts = turn.facts_to_remember or (None,)
            for fact in facts:
                if fact is None:
                    sheet.append([SCENARIO_NAMES[name], turn.turn_index, turn.transcript, "—", "", "", ""])
                else:
                    last = fact.valid_until or TURNS_PER_SCENARIO
                    sheet.append([
                        SCENARIO_NAMES[name],
                        turn.turn_index,
                        turn.transcript,
                        fact.text,
                        ", ".join(PLACE_NAMES[place] for place in fact.where),
                        "estar" if fact.expected else "haber desaparecido",
                        f"{turn.turn_index + 1} a {last}" if turn.turn_index < last else "—",
                    ])
                for cell in sheet[sheet.max_row]:
                    cell.alignment = _WRAP
    sheet.append([])
    sheet.append(["PDF adjunto (turno 1)", "", "", "", "", "", ""])
    sheet.cell(row=sheet.max_row, column=1).font = _TITLE_FONT
    for fact in PLANTED_FACTS:
        sheet.append(["PDF", 1, f"({fact.position} del documento) {fact.sentence}", fact.fact,
                      "respuesta", "estar", f"2 a {TURNS_PER_SCENARIO}"])
        for cell in sheet[sheet.max_row]:
            cell.alignment = _WRAP
    sheet.auto_filter.ref = f"A1:G{sheet.max_row}"


def write_excel(csv_path: Path, xlsx_path: Path | None = None) -> Path:
    """Lee el CSV del runner y escribe el Excel. Devuelve la ruta del Excel."""
    xlsx_path = xlsx_path or csv_path.with_suffix(".xlsx")
    with csv_path.open(encoding="utf-8", newline="") as file:
        rows = list(csv.DictReader(file))
    rows.sort(key=lambda r: (r["scenario"], int(r["attachment_kb"]), int(r["repeat"]), int(r["turn_index"])))

    workbook = Workbook()
    _turns_sheet(workbook, rows)
    _legend_sheet(workbook)
    _answer_key_sheet(workbook)
    workbook.save(xlsx_path)
    return xlsx_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Convierte results.csv del stress test en un Excel legible.")
    parser.add_argument("--csv", type=Path, required=True, help="CSV del runner (results.csv).")
    parser.add_argument("--output", type=Path, help="Excel de salida (por defecto, el CSV con .xlsx).")
    args = parser.parse_args(argv)
    print(f"Excel: {write_excel(args.csv, args.output)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
