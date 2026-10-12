"""Fabrica PDF y Word de prueba en memoria, para no guardar binarios en el repo.

No empieza por "test_", así que ni unittest ni pytest lo ejecutan como
pruebas: solo lo importan los tests que necesitan documentos.
"""

import io

from docx import Document
from pypdf import PdfWriter


def make_text_pdf(text: str) -> bytes:
    """PDF mínimo de una página con una línea de texto en Helvetica.

    Lo escribimos a mano, objeto a objeto, porque pypdf sabe leer texto pero
    no escribirlo. Un PDF es una lista de objetos numerados (catálogo,
    páginas, página, fuente, contenido) más una tabla final (xref) con la
    posición en bytes de cada uno.
    """
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_position = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_position,
    )
    return bytes(out)


def make_blank_pdf(pages: int = 1) -> bytes:
    """PDF de páginas en blanco: hace el papel de un PDF escaneado sin texto.

    Con ``pages`` > 1 sirve también para probar el tope de páginas.
    """
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_pdf_with_orphan_pages(texts: list[str]) -> bytes:
    """PDF trucado: páginas con texto, pero sin índice de páginas en el catálogo.

    pypdf no encuentra ninguna página (no hay árbol /Pages). pdfminer, el
    lector de MarkItDown, tiene un "modo rescate" que busca páginas sueltas
    por todo el fichero y SÍ las encuentra. Sirve para probar que los dos
    lectores no pueden ver documentos distintos.
    """
    objects = [
        b"<< /Type /Catalog >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for text in texts:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        objects.append(
            b"<< /Type /Page /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 2 0 R >> >> /Contents %d 0 R >>" % len(objects)
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_position = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_position,
    )
    return bytes(out)


def make_docx(title: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """Word con un título, unos párrafos y, si se pide, una tabla."""
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
