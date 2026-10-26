"""Tests del Excel legible del stress test (sesión 6, ``evals/stress/to_excel.py``).

Lanza una batería pequeña con el modelo de mentira (gratis), deja que el
runner escriba el CSV y el Excel, y comprueba que el Excel se lee bien: las
tres hojas, los títulos en castellano y que las columnas de "qué se ha
olvidado" lo cuentan con palabras.

Se ejecutan con:

    uv run python -m pytest tests/test_stress_excel.py -v
"""

from openpyxl import load_workbook

from evals.stress.run import RunConfig, main_async
from evals.stress.to_excel import COLUMNS, write_excel


async def run_fake_battery(tmp_path):
    config = RunConfig(
        scenarios=["contradiction", "pivot"], sizes_kb=[0, 5], repeats=1, turns=14, fake_llm=True,
        retry_wait_s=(0.0,), output=tmp_path / "results.csv", fixtures_dir=tmp_path,
    )
    await main_async(config)
    return config.output


def turns_by_key(workbook):
    sheet = workbook["Turnos"]
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    return headers, {(r["Escenario"], r["PDF (KB)"], r["Turno"]): r for r in rows}


async def test_runner_leaves_a_readable_excel_next_to_the_csv(tmp_path):
    csv_path = await run_fake_battery(tmp_path)
    xlsx_path = csv_path.with_suffix(".xlsx")
    assert xlsx_path.is_file()

    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["Turnos", "Leyenda", "Hoja de corrección"]
    headers, rows = turns_by_key(workbook)
    assert headers == [column.title for column in COLUMNS]
    assert len(rows) == 2 * 2 * 14  # 2 escenarios × 2 tamaños × 14 turnos

    # Lo que se olvida, con palabras: el 80.000 se pierde en el turno 14.
    assert "«80.000»" not in rows[("Contradicción", 0, 13)]["Olvidado o mal"]
    assert "«80.000»" in rows[("Contradicción", 0, 14)]["Olvidado o mal"]
    # Tras el pivote, React sigue donde no debería, y se dice dónde.
    pivot = rows[("Cambio de rumbo", 0, 6)]["Olvidado o mal"]
    assert "«react» sigue en la memoria" in pivot and "«react» sigue en la respuesta" in pivot
    # Los datos del PDF: los tres en el turno 1, ninguno en el 2.
    assert rows[("Contradicción", 5, 1)]["Datos del PDF en la respuesta"] == "✅ los 3"
    assert rows[("Contradicción", 5, 2)]["Datos del PDF en la respuesta"] == "❌ ninguno"
    assert rows[("Contradicción", 0, 1)]["Datos del PDF en la respuesta"] == "sin PDF"
    assert rows[("Contradicción", 0, 1)]["Resultado"] == "✅"


async def test_legend_explains_every_column_and_the_answer_key_lists_the_scripts(tmp_path):
    csv_path = await run_fake_battery(tmp_path)
    workbook = load_workbook(write_excel(csv_path, tmp_path / "otro_nombre.xlsx"))

    legend = [row[0] for row in workbook["Leyenda"].iter_rows(min_row=2, values_only=True)]
    assert all(column.title in legend for column in COLUMNS)

    answer_key = list(workbook["Hoja de corrección"].iter_rows(min_row=2, values_only=True))
    assert ("Contradicción", 3) in {(row[0], row[1]) for row in answer_key}
    assert any(row[3] == "80.000" and row[6] == "9 a 20" for row in answer_key)
    assert any(row[0] == "PDF" and row[3] == "12.000 reservas" for row in answer_key)
